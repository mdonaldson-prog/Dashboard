"""Aurora cycle count schedule: velocity from inFlow Sales Orders + Product Details + BOM exports.
Google Sheets compatible (no dynamic arrays, no circular references).
Usage: python3 build_count_sheet.py out.xlsx [exports_folder]
With a folder, the inFlow SalesOrder, ProductDetails and BOM CSVs in it are pre-loaded.
That copy holds business data: keep it out of git."""
import sys, csv, glob
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

out = sys.argv[1]
with_data = len(sys.argv) > 2
U = sys.argv[2] if with_data else ""
def find(part):
    hits = sorted(glob.glob(f"{U}/*{part}*.csv"))
    if not hits: raise SystemExit(f"No *{part}*.csv in {U}")
    return hits[-1]

SO_N, PR_N, BOM_N = 30000, 3000, 5000      # data rows supported
SO_L, PR_L, BOM_L = SO_N + 1, PR_N + 1, BOM_N + 1
LEVELS = 6                                  # BOM levels rolled up (data has 5)
WEEKS, DAILY_ROWS = 13, 200

F = "Arial"
font = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=14, bold=True); blue = Font(name=F, size=10, color="0000FF")
note = Font(name=F, size=9, italic=True, color="808080"); white_b = Font(name=F, size=10, bold=True, color="FFFFFF")
warn = Font(name=F, size=10, bold=True, color="C00000")
hdr_fill = PatternFill("solid", fgColor="1F3864"); paste_hdr = PatternFill("solid", fgColor="2E75B6")
input_fill = PatternFill("solid", fgColor="FFF2CC"); calc_fill = PatternFill("solid", fgColor="F2F2F2")
cls_fill = {"A": "C6EFCE", "B": "FFEB9C", "C": "F8CBAD", "D": "D9D9D9"}
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")

wb = Workbook()

def hdr(ws, row, labels, fill=hdr_fill, start=1):
    for i, l in enumerate(labels, start):
        c = ws.cell(row=row, column=i, value=l)
        c.font, c.fill, c.border = white_b, fill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

def cls_cf(ws, rng, first_cell):
    for k, col in cls_fill.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first_cell}="{k}"'], fill=PatternFill("solid", fgColor=col)))

# ======================================================================= How to use
hw = wb.active; hw.title = "How to use"
lines = [
    ("Aurora Cycle Count Schedule", title),
    ("Velocity comes from inFlow sales orders. A SKUs are counted weekly, B every 2 weeks, C every 4 weeks, and D (no sales in 90 days) once a year. Counting days are Tuesday to Friday.", font),
    ("", font),
    ("Once a month: refresh the data", bold),
    ("1. In inFlow, export Sales Orders with order dates covering at least the last 90 days (inFlow_SalesOrder.csv).", font),
    ("2. Go to the 'Sales Orders' tab. Select all (click the box left of column A), press Delete, click cell A1, then File → Import → Upload → choose the CSV → Import location 'Replace data at selected cell'.", font),
    ("3. Only when products or BOMs have changed: do the same with inFlow Product Details on the 'Products' tab and the inFlow BOM on the 'BOM' tab.", font),
    ("4. On Settings, set 'Schedule start' to the Tuesday you want the new 4-week cycle to begin, and check the capacity line.", font),
    ("5. Check the Velocity tab (SKUs ranked by transactions, with their class and count day) and the Calendar tab (counts per day against capacity).", font),
    ("", font),
    ("Every counting day", bold),
    ("Open 'Daily Count Sheet', pick the date in cell B3, and print. System quantity isn't shown (blind count). Record the counts in inFlow's Stock Count as usual.", font),
    ("", font),
    ("How velocity is worked out", bold),
    ("Transactions = sales order lines for the SKU over the last 90 days (to the latest order date in the file), at Aurora plus lines with no location (switchable on Settings). Quotes and cancelled orders are left out.", font),
    ("BOM: each sales line for a kit or finished product also counts as one transaction for every component in its BOM, down through all BOM levels, so raw materials used to make what sells are counted too. Settings can limit this to auto-manufactured kits, or switch it off.", font),
    ("SKUs are ranked by transactions. A = the SKUs making up the first 80% of transactions, B = the next 15%, C = the rest that moved, D = stocked products with no transactions.", font),
    ("Only active, stocked products from the Products tab are scheduled. Non-stocked products and services are skipped.", font),
    ("", font),
    ("Overrides", bold),
    ("On the Overrides tab, type a SKU and a class (A, B, C or D) to pin it, or X to stop counting it (e.g. a product that isn't kept at Aurora). Overrides follow the SKU, so they survive each month's data refresh.", font),
    ("", font),
    ("How SKUs are spread over the days", bold),
    ("Within each class, SKUs are dealt round-robin in velocity order across the counting days of one cycle (4 days for A, 8 for B, 16 for C, 208 for D), so each day gets an even share. A SKU keeps its day until the classes change at the next monthly refresh.", font),
    ("Holidays aren't skipped automatically: count that day's list on the nearest open counting day.", font),
    ("", font),
    ("Cell colors", bold),
    ("Yellow: settings and overrides you can edit.   Blue headers: tabs you paste inFlow exports into.   Grey: formulas, don't edit.", font),
    ("", font),
    ("Example override (Overrides tab)", bold),
]
for i, (t, f) in enumerate(lines, 1):
    hw.cell(row=i, column=1, value=t).font = f
r = len(lines) + 1
hdr(hw, r, ["SKU", "Class", "Reason"], input_fill)
for j in range(1, 4): hw.cell(row=r, column=j).font = bold
for j, val in enumerate(["50009", "A", "High-value flake, count weekly"], 1):
    c = hw.cell(row=r + 1, column=j, value=val); c.font, c.border = font, box
hw.cell(row=r + 2, column=1, value="(Example only, shown here so it doesn't change the schedule.)").font = note
widths(hw, {"A": 14, "B": 10, "C": 34})

# ======================================================================= Settings
st = wb.create_sheet("Settings")
st["A1"] = "Settings"; st["A1"].font = title
inputs = [
    (3, "Location", "Aurora", None, "Must match the Location in inFlow exactly."),
    (4, "Count sales lines with no location", "Yes", None, "Yes = treat lines with a blank Location as this location (as the Dashboard does)."),
    (5, "Velocity window (days)", 90, "0", "Transactions are counted over this many days, ending on the latest order date in the file."),
    (6, "A = first % of transactions", 0.80, "0%", "SKUs that make up the top 80% of transactions."),
    (7, "B = up to % of transactions", 0.95, "0%", "The next 15%. The rest that moved are C."),
    (8, "BOM components", "All BOMs", None, "All BOMs / Kits only (auto-manufactured) / Off. See How to use."),
    (9, "Schedule start (a Tuesday)", date(2026, 10, 20), "ddd m/d/yyyy", "Week 1 of the cycle starts here."),
    (10, "Daily capacity (SKUs per day)", 50, "0", "Confirmed range is 40–60."),
    (11, "Schedule D SKUs", "Yes", None, "Yes = count D SKUs on the D cadence below. No = leave them off."),
]
for rr, label, val, fmt, nt in inputs:
    st.cell(row=rr, column=1, value=label).font = bold
    c = st.cell(row=rr, column=2, value=val); c.font, c.fill, c.border = blue, input_fill, box
    if fmt: c.number_format = fmt
    st.cell(row=rr, column=3, value=nt).font = note
for cell, opts in (("B4", "Yes,No"), ("B11", "Yes,No"), ("B8", "All BOMs,Kits only,Off")):
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=False); st.add_data_validation(dv); dv.add(cell)
st["D9"] = '=IF(WEEKDAY(B9)=3,"","⚠ Not a Tuesday")'; st["D9"].font = warn

st["A13"] = "Data loaded"; st["A13"].font = bold
st["A14"] = "Sales lines in file"; st["B14"] = f"=COUNTIF('SO Calc'!$A$2:$A${SO_L},\"?*\")"
st["A15"] = "Latest order date"; st["B15"] = f"=IF(COUNT('SO Calc'!$C$2:$C${SO_L})=0,\"\",MAX('SO Calc'!$C$2:$C${SO_L}))"
st["A16"] = "Velocity window starts after"; st["B16"] = '=IF(B15="","",B15-B5)'
st["A17"] = "Sales lines used (in window, this location)"; st["B17"] = f"=COUNTIF('SO Calc'!$D$2:$D${SO_L},\"?*\")"
st["A18"] = "Stocked products"; st["B18"] = f"=COUNT('Product Calc'!$G$2:$G${PR_L})"
st["A19"] = "Earliest order date in file"; st["B19"] = f"=IF(COUNT('SO Calc'!$C$2:$C${SO_L})=0,\"\",MIN('SO Calc'!$C$2:$C${SO_L}))"
st["C19"] = '=IF(B19="","Paste the Sales Orders export",IF(B19>B16+1,"⚠ File starts "&TEXT(B19,"m/d/yyyy")&": less than the velocity window, so velocity is understated",""))'
st["C19"].font = warn
for c in ("B15", "B16", "B19"): st[c].number_format = "m/d/yyyy"
for rr in range(14, 20):
    st.cell(row=rr, column=1).font = font; st.cell(row=rr, column=2).font = bold; st.cell(row=rr, column=2).fill = calc_fill

# export column positions, found by header name
st["A21"] = "Export columns (found automatically)"; st["A21"].font = bold
cols = [(22, "Sales Orders", "OrderDate"), (23, "Sales Orders", "Location"), (24, "Sales Orders", "ProductSKU"),
        (25, "Sales Orders", "IsQuote"), (26, "Sales Orders", "IsCancelled"),
        (27, "Products", "SKU"), (28, "Products", "ProductName"), (29, "Products", "Category"), (30, "Products", "ItemType"),
        (31, "Products", "AutoManufacture"), (32, "Products", "IsActive"),
        (33, "BOM", "FinishedProductSKU"), (34, "BOM", "ComponentProductSKU"), (35, "BOM", "IsActive")]
POS = {}
for rr, tab, name in cols:
    st.cell(row=rr, column=1, value=f"{tab}: {name}").font = font
    st.cell(row=rr, column=2, value=f"=IFERROR(MATCH(\"{name}\",'{tab}'!$1:$1,0),\"missing\")").font = font
    st.cell(row=rr, column=2).fill = calc_fill
    POS[(tab, name)] = f"Settings!$B${rr}"
st["C22"] = '=IF(COUNTIF(B22:B35,"missing")>0,"⚠ A column is missing: check the export headers on row 1 of each data tab","")'
st["C22"].font = warn

# class table
CT = 38
for j, l in enumerate(["Class", "Rule", "Cadence (weeks)", "Counting days per cycle", "SKUs", "SKUs per day", "Share of transactions"], 1):
    c = st.cell(row=CT, column=j, value=l); c.font, c.fill, c.border = white_b, hdr_fill, box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
rules = [("A", "First 80%", 1), ("B", "Next 15%", 2), ("C", "Rest that moved", 4), ("D", "No transactions", 52)]
for i, (k, rule, cad) in enumerate(rules):
    rr = CT + 1 + i
    st.cell(row=rr, column=1, value=k).font = bold
    st.cell(row=rr, column=1).fill = PatternFill("solid", fgColor=cls_fill[k])
    st.cell(row=rr, column=2, value=rule).font = font
    c = st.cell(row=rr, column=3, value=cad); c.font, c.fill = blue, input_fill
    st.cell(row=rr, column=4, value=f"=C{rr}*4")
    st.cell(row=rr, column=5, value=f"=COUNTIF('Product Calc'!$O$2:$O${PR_L},A{rr})")
    act = f'OR(A{rr}<>"D",$B$11="Yes")'
    st.cell(row=rr, column=6, value=f"=IF({act},E{rr}/D{rr},0)").number_format = "0.0"
    st.cell(row=rr, column=7, value=f"=IF($B$43=0,0,SUMIFS('Product Calc'!$J$2:$J${PR_L},'Product Calc'!$O$2:$O${PR_L},A{rr})/$B$43)").number_format = "0%"
    for j in range(1, 8):
        st.cell(row=rr, column=j).border = box
        if j >= 4: st.cell(row=rr, column=j).fill = calc_fill; st.cell(row=rr, column=j).font = font
TR = CT + 5  # 43
st.cell(row=TR, column=1, value="Total").font = bold
st.cell(row=TR, column=2, value=f"=SUMIFS('Product Calc'!$J$2:$J${PR_L},'Product Calc'!$G$2:$G${PR_L},1)").number_format = '#,##0" tx"'
st.cell(row=TR, column=5, value=f"=SUM(E{CT+1}:E{TR-1})")
st.cell(row=TR, column=6, value=f"=SUM(F{CT+1}:F{TR-1})").number_format = "0.0"
st.cell(row=TR, column=7, value=f"=SUM(G{CT+1}:G{TR-1})").number_format = "0%"
for j in range(1, 8): st.cell(row=TR, column=j).border = box; st.cell(row=TR, column=j).font = bold
st.cell(row=TR + 1, column=1, value="Capacity check").font = bold
st.cell(row=TR + 1, column=2, value=f'=IF(E{TR}=0,"Paste the exports on the Sales Orders, Products and BOM tabs",IF(F{TR}<=$B$10,"OK: "&TEXT(F{TR},"0.0")&" SKUs a day needed, capacity "&$B$10,"Over capacity by "&TEXT(F{TR}-$B$10,"0.0")&" SKUs a day. Options: raise capacity, lower the A %, lengthen the C or D cadence, or mark SKUs not kept here as X on Overrides"))').font = bold
st.cell(row=TR + 2, column=1, value="Total transactions (B43) are for stocked products only. 'SKUs per day' is the average load: SKUs ÷ (cadence weeks × 4 counting days).").font = note
widths(st, {"A": 38, "B": 16, "C": 16, "D": 13, "E": 10, "F": 11, "G": 13})

S_LOC, S_BLANK, S_DAYS, S_CUTA, S_CUTB, S_BOM = "Settings!$B$3", "Settings!$B$4", "Settings!$B$5", "Settings!$B$6", "Settings!$B$7", "Settings!$B$8"
S_START, S_CAP, S_DORM, S_END = "Settings!$B$9", "Settings!$B$10", "Settings!$B$11", "Settings!$B$15"
S_TOT = "Settings!$B$43"
CAD, CLS = f"Settings!$C${CT+1}:$C${CT+4}", f"Settings!$A${CT+1}:$A${CT+4}"

# ======================================================================= Overrides
ov = wb.create_sheet("Overrides")
ov["A1"] = "Overrides"; ov["A1"].font = title
ov["A2"] = "Pin a SKU to a class (A, B, C or D), or X to stop counting it. Columns D–E fill in by themselves."; ov["A2"].font = note
hdr(ov, 4, ["SKU", "Class", "Reason", "Product", "Calculated class"])
for rr in range(5, 505):
    for col in "ABC":
        ov[f"{col}{rr}"].fill = input_fill; ov[f"{col}{rr}"].font = blue
    m = f"MATCH(TRIM(A{rr}&\"\"),'Product Calc'!$A$2:$A${PR_L},0)"
    ov[f"D{rr}"] = f'=IF(A{rr}="","",IFERROR(INDEX(\'Product Calc\'!$B$2:$B${PR_L},{m}),"⚠ SKU not found"))'
    ov[f"E{rr}"] = f'=IF(A{rr}="","",IFERROR(INDEX(\'Product Calc\'!$M$2:$M${PR_L},{m}),""))'
    ov[f"F{rr}"] = f'=TRIM(A{rr}&"")'
    for col in "DE": ov[f"{col}{rr}"].font = font; ov[f"{col}{rr}"].fill = calc_fill
dvo = DataValidation(type="list", formula1='"A,B,C,D,X"', allow_blank=True); ov.add_data_validation(dvo); dvo.add("B5:B504")
ov.column_dimensions["F"].hidden = True
widths(ov, {"A": 16, "B": 8, "C": 36, "D": 44, "E": 11})
ov.freeze_panes = "A5"

# ======================================================================= Velocity (ranked view)
ve = wb.create_sheet("Velocity")
VH = ["Rank", "SKU", "Product", "Category", "Sales lines", "BOM lines", "Transactions", "Cumulative %", "Calculated class", "Override", "Class", "Count day", "Cycle week", "Next count"]
hdr(ve, 1, VH)
src = {"B": "A", "C": "B", "D": "C", "E": "H", "F": "I", "G": "J", "I": "M", "J": "N", "K": "O", "L": "U", "M": "V", "N": "X"}
for rr in range(2, PR_L + 1):
    n = rr - 1
    ve[f"A{rr}"] = f"=IF({n}>COUNT('Product Calc'!$K$2:$K${PR_L}),\"\",{n})"
    m = f"MATCH($A{rr},'Product Calc'!$K$2:$K${PR_L},0)"
    for col, pc in src.items():
        ve[f"{col}{rr}"] = f"=IF($A{rr}=\"\",\"\",INDEX('Product Calc'!${pc}$2:${pc}${PR_L},{m}))"
    ve[f"H{rr}"] = f"=IF($A{rr}=\"\",\"\",IF(G{rr}=0,\"\",INDEX('Product Calc'!$L$2:$L${PR_L},{m})+G{rr}/{S_TOT}))"
    ve[f"H{rr}"].number_format = "0.0%"; ve[f"N{rr}"].number_format = "ddd m/d/yy"
    for col in "ABCDEFGHIJKLMN": ve[f"{col}{rr}"].font = font
cls_cf(ve, f"K2:K{PR_L}", "$K2"); cls_cf(ve, f"I2:I{PR_L}", "$I2")
widths(ve, {"A": 6, "B": 14, "C": 42, "D": 16, "E": 8, "F": 8, "G": 11, "H": 11, "I": 10, "J": 9, "K": 7, "L": 9, "M": 11, "N": 12})
ve.freeze_panes = "C2"

# ======================================================================= Calendar
cal = wb.create_sheet("Calendar")
cal["A1"] = "Count Calendar"; cal["A1"].font = title
cal["A2"] = f'={S_LOC}&": {WEEKS} weeks from "&TEXT({S_START},"ddd m/d/yyyy")&". Capacity "&{S_CAP}&" SKUs a day."'
cal["A2"].font = font
hdr(cal, 4, ["Date", "Day", "Week", "A (weekly)", "B (2 weeks)", "C (4 weeks)", "D (yearly)", "Total", "Capacity", "Room left"])
first = 5
for i in range(WEEKS * 4):
    rr = first + i
    cal[f"A{rr}"] = f"={S_START}+INT(({rr}-{first})/4)*7+MOD({rr}-{first},4)"; cal[f"A{rr}"].number_format = "m/d/yyyy"
    cal[f"B{rr}"] = f'=CHOOSE(MOD({rr}-{first},4)+1,"Tue","Wed","Thu","Fri")'
    cal[f"C{rr}"] = f"=INT(({rr}-{first})/4)+1"
    for j, (k, colL) in enumerate(zip("ABCD", "DEFG")):
        cad = f"Settings!$C${CT+1+j}"
        cal[f"{colL}{rr}"] = (f"=COUNTIFS('Product Calc'!$O$2:$O${PR_L},\"{k}\",'Product Calc'!$S$2:$S${PR_L},MOD({rr}-{first},4),"
                              f"'Product Calc'!$T$2:$T${PR_L},MOD($C{rr}-1,{cad}))")
    cal[f"H{rr}"] = f"=SUM(D{rr}:G{rr})"; cal[f"I{rr}"] = f"={S_CAP}"; cal[f"J{rr}"] = f"=I{rr}-H{rr}"
    for col in "ABCDEFGHIJ":
        c = cal[f"{col}{rr}"]; c.font, c.border = font, box
        if col != "A": c.alignment = center
        if (i // 4) % 2: c.fill = PatternFill("solid", fgColor="F7F9FC")
    cal[f"H{rr}"].font = bold
last_cal = first + WEEKS * 4 - 1
cal.conditional_formatting.add(f"H{first}:H{last_cal}", FormulaRule(formula=[f"$H{first}>$I{first}"], fill=PatternFill("solid", fgColor="F4B6B6"), font=Font(name=F, bold=True, color="9C0006")))
cal.conditional_formatting.add(f"J{first}:J{last_cal}", FormulaRule(formula=[f"$J{first}<0"], font=Font(name=F, color="C00000")))
tr = last_cal + 1
cal[f"A{tr}"] = "Total"; cal[f"A{tr+1}"] = "Average per day"
for col in "DEFGH":
    cal[f"{col}{tr}"] = f"=SUM({col}{first}:{col}{last_cal})"
    cal[f"{col}{tr+1}"] = f"=AVERAGE({col}{first}:{col}{last_cal})"; cal[f"{col}{tr+1}"].number_format = "0.0"
    for x in (tr, tr + 1): cal[f"{col}{x}"].alignment = center; cal[f"{col}{x}"].font = bold
cal[f"A{tr}"].font = cal[f"A{tr+1}"].font = bold
cal[f"A{tr+3}"] = "To plan further ahead, select the last 4 rows and drag them down."; cal[f"A{tr+3}"].font = note
widths(cal, {"A": 12, "B": 7, "C": 7, "D": 11, "E": 11, "F": 11, "G": 11, "H": 9, "I": 10, "J": 10})
cal.freeze_panes = "A5"

# ======================================================================= Daily Count Sheet
ds = wb.create_sheet("Daily Count Sheet")
ds["A1"] = f'="Daily Cycle Count Sheet: "&{S_LOC}'; ds["A1"].font = title
ds["A3"] = "Count date"; ds["A3"].font = bold
ds["B3"] = f"=Calendar!A{first}"
ds["B3"].number_format = "dddd, mmmm d, yyyy"; ds["B3"].font = Font(name=F, size=11, bold=True, color="0000FF"); ds["B3"].fill = input_fill; ds["B3"].border = box
ds.merge_cells("B3:C3")
dvd = DataValidation(type="list", formula1=f"=Calendar!$A${first}:$A${last_cal}", allow_blank=False); ds.add_data_validation(dvd); dvd.add("B3")
ds["B3"].comment = Comment("Pick a date from the list, or type any Tuesday to Friday on or after the schedule start.", "Schedule")
ds["D3"] = f'=IF(OR(J3<0,J3>3),"⚠ Not a counting day in this schedule","")'; ds["D3"].font = warn
ds["J3"] = f"=IF(B3<{S_START},-1,MOD(B3-{S_START},7))"; ds["K3"] = f"=INT((B3-{S_START})/7)"
for c in ("J3", "K3"): ds[c].font = Font(name=F, size=8, color="A6A6A6")
ds["A4"] = "SKUs to count"; ds["A4"].font = bold
ds["B4"] = f"=COUNT('Product Calc'!$Y$2:$Y${PR_L})"; ds["B4"].font = bold; ds["B4"].alignment = Alignment(horizontal="left")
ds["C4"] = "=" + '&" · "&'.join(f"\"{k} \"&COUNTIFS('Product Calc'!$Y$2:$Y${PR_L},1,'Product Calc'!$O$2:$O${PR_L},\"{k}\")" for k in "ABCD")
ds["C4"].font = font
ds["A5"] = "Counted by: ____________________    Blind count: system quantity isn't shown. Count A SKUs before picking starts, and every bin for each SKU."
ds["A5"].font = Font(name=F, size=9, italic=True)
hdr(ds, 7, ["#", "SKU", "Product", "Category", "Class", "Counted qty", "UoM / notes"])
for n in range(1, DAILY_ROWS + 1):
    rr = 7 + n
    m = f"MATCH($A{rr},'Product Calc'!$AA$2:$AA${PR_L},0)"
    ds[f"A{rr}"] = f'=IF({n}<=$B$4,{n},"")'
    for col, pc in (("B", "A"), ("C", "B"), ("D", "C"), ("E", "O")):
        ds[f"{col}{rr}"] = f"=IF($A{rr}=\"\",\"\",INDEX('Product Calc'!${pc}$2:${pc}${PR_L},{m}))"
    for col in "ABCDEFG":
        c = ds[f"{col}{rr}"]; c.font, c.border = font, box
    ds[f"A{rr}"].alignment = ds[f"E{rr}"].alignment = center
cls_cf(ds, f"E8:E{7+DAILY_ROWS}", "$E8")
widths(ds, {"A": 5, "B": 15, "C": 44, "D": 16, "E": 7, "F": 12, "G": 22, "J": 4, "K": 4})
ds.freeze_panes = "A8"; ds.print_title_rows = "7:7"; ds.print_area = f"A1:G{7+DAILY_ROWS}"
ds.page_setup.orientation = "portrait"; ds.page_setup.fitToWidth = 1; ds.page_setup.fitToHeight = 0
ds.sheet_properties.pageSetUpPr.fitToPage = True

# ======================================================================= raw paste tabs
so_t = wb.create_sheet("Sales Orders"); pr_t = wb.create_sheet("Products"); bom_t = wb.create_sheet("BOM")
for t in (so_t, pr_t, bom_t): t.sheet_properties.tabColor = "2E75B6"

def load(ws, path, keep=None):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    head = rows[0]
    idx = [head.index(k) for k in keep] if keep else list(range(len(head)))
    for j, i in enumerate(idx, 1):
        c = ws.cell(row=1, column=j, value=head[i]); c.font, c.fill = white_b, paste_hdr
    for r_i, row in enumerate(rows[1:], 2):
        for j, i in enumerate(idx, 1):
            v = row[i]
            if v != "": ws.cell(row=r_i, column=j, value=v)
    return len(rows) - 1

SO_KEEP = ["OrderNumber", "InventoryStatus", "OrderDate", "Location", "ProductName", "ProductSKU", "ProductQuantity", "IsQuote", "IsCancelled"]
if with_data:
    n_so = load(so_t, find("SalesOrder"), SO_KEEP)
    n_pr = load(pr_t, find("ProductDetails"))
    n_bom = load(bom_t, find("BOM"))
    assert n_so <= SO_N and n_pr <= PR_N and n_bom <= BOM_N, (n_so, n_pr, n_bom)
    print("loaded", n_so, n_pr, n_bom)
else:
    for ws, heads in ((so_t, SO_KEEP), (pr_t, ["ProductName", "SKU", "Category", "ItemType", "AutoManufacture", "IsActive"]),
                      (bom_t, ["FinishedProduct", "FinishedProductSKU", "ComponentProduct", "ComponentProductSKU", "Quantity", "QuantityUom", "IsActive"])):
        for j, h in enumerate(heads, 1):
            c = ws.cell(row=1, column=j, value=h); c.font, c.fill = white_b, paste_hdr
    so_t["K1"] = "← Paste the full inFlow export here (replace row 1 too). Columns are found by name, so extra columns are fine."
    so_t["K1"].font = note

def bool_(x):  # TRUE however the export lands: TRUE, "True", 1
    return f'OR(UPPER(TRIM({x}&""))="TRUE",{x}&""="1")'

# ======================================================================= SO Calc
sc = wb.create_sheet("SO Calc")
hdr(sc, 1, ["SKU", "Date text", "Order date", "Counts toward velocity (SKU)"])
SO = f"'Sales Orders'!$A$1:$CZ${SO_L}"
P = lambda name: POS[("Sales Orders", name)]
for rr in range(2, SO_L + 1):
    raw = lambda name: f"INDEX({SO},{rr},{P(name)})"
    sc[f"A{rr}"] = f'=IFERROR(TRIM({raw("ProductSKU")}&""),"")'
    sc[f"B{rr}"] = f'=IF(A{rr}="","",IFERROR(LEFT({raw("OrderDate")}&" ",FIND(" ",{raw("OrderDate")}&" ")-1),""))'
    b = f"B{rr}"; p1 = f'FIND("/",{b})'; p2 = f'FIND("/",{b},{p1}+1)'
    sc[f"C{rr}"] = (f'=IF(A{rr}="","",IFERROR(IF(ISNUMBER({raw("OrderDate")}),INT({raw("OrderDate")}),'
                    f'DATE(VALUE(MID({b},{p2}+1,4)),VALUE(LEFT({b},{p1}-1)),VALUE(MID({b},{p1}+1,{p2}-{p1}-1)))),""))')
    loc = f'TRIM(IFERROR({raw("Location")},"")&"")'
    sc[f"D{rr}"] = (f'=IF(OR(A{rr}="",C{rr}=""),"",IF(AND(C{rr}>{S_END}-{S_DAYS},OR({loc}={S_LOC},AND({loc}="",{S_BLANK}="Yes")),'
                    f'NOT({bool_(f"IFERROR({raw("IsQuote")},FALSE)")}),NOT({bool_(f"IFERROR({raw("IsCancelled")},FALSE)")})),A{rr},""))')
    sc[f"C{rr}"].number_format = "m/d/yyyy"
widths(sc, {"A": 14, "B": 12, "C": 12, "D": 16})
sc.sheet_properties.tabColor = "A6A6A6"

# ======================================================================= Product Calc
pc = wb.create_sheet("Product Calc")
lvl_cols = []  # rollup level columns: AB.. (level 1..LEVELS)
from openpyxl.utils import get_column_letter as gcl
base = 28  # AB
for k in range(1, LEVELS + 1): lvl_cols.append(gcl(base + k - 1))
heads = ["SKU", "Product", "Category", "Item type", "Auto-manufacture", "First row for SKU", "Stocked (classified)",
         "Sales lines", "BOM lines", "Transactions", "Rank", "Cumulative % before", "Calculated class", "Override", "Class",
         "Cadence (weeks)", "# in class", "_slot", "_day (0=Tue)", "_cycle week (0-based)", "Count day", "Cycle week",
         "First count", "Next count", "Due on daily-sheet date", "Daily sort key", "Daily #"] + [f"Tx level {k}" for k in range(1, LEVELS + 1)]
hdr(pc, 1, heads)
PR = f"Products!$A$1:$BZ${PR_L}"
Q = lambda name: POS[("Products", name)]
L = PR_L
for rr in range(2, PR_L + 1):
    raw = lambda name: f"INDEX({PR},{rr},{Q(name)})"
    pc[f"A{rr}"] = f'=IFERROR(TRIM({raw("SKU")}&""),"")'
    pc[f"B{rr}"] = f'=IF(A{rr}="","",IFERROR({raw("ProductName")}&"",""))'
    pc[f"C{rr}"] = f'=IF(A{rr}="","",IFERROR({raw("Category")}&"",""))'
    pc[f"D{rr}"] = f'=IF(A{rr}="","",IFERROR({raw("ItemType")}&"",""))'
    pc[f"E{rr}"] = f'=IF(A{rr}="","",IFERROR({bool_(raw("AutoManufacture"))},FALSE))'
    pc[f"F{rr}"] = f'=IF(A{rr}="","",IF(COUNTIF($A$2:A{rr},A{rr})=1,1,""))'
    active = f'IFERROR(NOT(UPPER(TRIM({raw("IsActive")}&""))="FALSE"),TRUE)'
    pc[f"G{rr}"] = f'=IF(F{rr}<>1,"",IF(AND(D{rr}="Stocked product",{active}),1,""))'
    pc[f"H{rr}"] = f"=IF(F{rr}<>1,\"\",COUNTIF('SO Calc'!$D$2:$D${SO_L},A{rr}))"
    pc[f"J{rr}"] = f"=IF(F{rr}<>1,\"\",{lvl_cols[-1]}{rr})"
    pc[f"I{rr}"] = f'=IF(F{rr}<>1,"",J{rr}-H{rr})'
    for k, col in enumerate(lvl_cols, 1):
        prev = f"H{rr}" if k == 1 else f"{lvl_cols[k-2]}{rr}"
        bcol = gcl(6 + k)  # BOM Calc G.. (level k parent tx)
        pc[f"{col}{rr}"] = f"=IF(F{rr}<>1,\"\",H{rr}+SUMIFS('BOM Calc'!${bcol}$2:${bcol}${BOM_L},'BOM Calc'!$B$2:$B${BOM_L},A{rr}))"
    G, J = f"$G$2:$G${L}", f"$J$2:$J${L}"
    pc[f"K{rr}"] = f'=IF(G{rr}<>1,"",COUNTIFS({G},1,{J},">"&J{rr})+COUNTIFS($G$2:G{rr},1,$J$2:J{rr},J{rr}))'
    pc[f"L{rr}"] = f'=IF(G{rr}<>1,"",IF({S_TOT}=0,1,(SUMIFS({J},{G},1,{J},">"&J{rr})+J{rr}*(COUNTIFS($G$2:G{rr},1,$J$2:J{rr},J{rr})-1))/{S_TOT}))'
    pc[f"M{rr}"] = f'=IF(G{rr}<>1,"",IF(J{rr}=0,"D",IF(L{rr}<{S_CUTA},"A",IF(L{rr}<{S_CUTB},"B","C"))))'
    pc[f"N{rr}"] = f'=IF(G{rr}<>1,"",IFERROR(UPPER(TRIM(INDEX(Overrides!$B$5:$B$504,MATCH(A{rr},Overrides!$F$5:$F$504,0))&"")),""))'
    pc[f"O{rr}"] = f'=IF(G{rr}<>1,"",IF(OR(N{rr}="A",N{rr}="B",N{rr}="C",N{rr}="D",N{rr}="X"),N{rr},M{rr}))'
    pc[f"P{rr}"] = f'=IF(OR(O{rr}="",O{rr}="X"),"",IF(AND(O{rr}="D",{S_DORM}<>"Yes"),"",IFERROR(INDEX({CAD},MATCH(O{rr},{CLS},0)),"")))'
    pc[f"Q{rr}"] = f'=IF(P{rr}="","",COUNTIFS($O$2:$O${L},O{rr},$K$2:$K${L},"<="&K{rr}))'
    pc[f"R{rr}"] = f'=IF(P{rr}="","",MOD(Q{rr}-1,P{rr}*4))'
    pc[f"S{rr}"] = f'=IF(P{rr}="","",MOD(R{rr},4))'
    pc[f"T{rr}"] = f'=IF(P{rr}="","",INT(R{rr}/4))'
    pc[f"U{rr}"] = f'=IF(P{rr}="","",CHOOSE(S{rr}+1,"Tue","Wed","Thu","Fri"))'
    pc[f"V{rr}"] = f'=IF(P{rr}="",IF(O{rr}="X","Not counted",IF(O{rr}="D","Off","")),IF(P{rr}=1,"Every week",T{rr}+1&" of "&P{rr}))'
    pc[f"W{rr}"] = f'=IF(P{rr}="","",{S_START}+T{rr}*7+S{rr})'
    pc[f"X{rr}"] = f'=IF(P{rr}="","",W{rr}+MAX(0,ROUNDUP((TODAY()-W{rr})/(P{rr}*7),0))*P{rr}*7)'
    pc[f"Y{rr}"] = f"=IF(P{rr}=\"\",\"\",IF(AND('Daily Count Sheet'!$J$3=S{rr},MOD('Daily Count Sheet'!$K$3,P{rr})=T{rr}),1,\"\"))"
    pc[f"Z{rr}"] = f'=IF(Y{rr}=1,FIND(O{rr},"ABCD")*100000+K{rr},"")'
    pc[f"AA{rr}"] = f'=IF(Y{rr}=1,COUNTIFS($Y$2:$Y${L},1,$Z$2:$Z${L},"<="&Z{rr}),"")'
    pc[f"L{rr}"].number_format = "0.0%"
    for c in ("W", "X"): pc[f"{c}{rr}"].number_format = "m/d/yyyy"
widths(pc, {"A": 14, "B": 30})
pc.freeze_panes = "B2"
pc.sheet_properties.tabColor = "A6A6A6"

# ======================================================================= BOM Calc
bc = wb.create_sheet("BOM Calc")
hdr(bc, 1, ["Finished SKU", "Component SKU", "Active", "Finished is auto-manufactured", "Used (per BOM setting)", "Finished row in Product Calc"]
    + [f"Finished tx, level {k}" for k in range(1, LEVELS + 1)])
BM = f"BOM!$A$1:$Z${BOM_L}"
Bp = lambda name: POS[("BOM", name)]
for rr in range(2, BOM_L + 1):
    raw = lambda name: f"INDEX({BM},{rr},{Bp(name)})"
    bc[f"A{rr}"] = f'=IFERROR(TRIM({raw("FinishedProductSKU")}&""),"")'
    bc[f"B{rr}"] = f'=IFERROR(TRIM({raw("ComponentProductSKU")}&""),"")'
    bc[f"C{rr}"] = f'=IF(A{rr}="","",IFERROR(NOT(UPPER(TRIM({raw("IsActive")}&""))="FALSE"),TRUE))'
    bc[f"F{rr}"] = f"=IF(OR(A{rr}=\"\",B{rr}=\"\"),\"\",IFERROR(MATCH(A{rr},'Product Calc'!$A$2:$A${PR_L},0),\"\"))"
    bc[f"D{rr}"] = f"=IF(F{rr}=\"\",FALSE,INDEX('Product Calc'!$E$2:$E${PR_L},F{rr})=TRUE)"
    bc[f"E{rr}"] = f'=IF(F{rr}="",0,IF(AND(C{rr},OR({S_BOM}="All BOMs",AND({S_BOM}="Kits only",D{rr}))),1,0))'
    for k in range(1, LEVELS + 1):
        col = gcl(6 + k)
        src_col = "H" if k == 1 else lvl_cols[k - 2]
        bc[f"{col}{rr}"] = f"=IF(E{rr}=1,INDEX('Product Calc'!${src_col}$2:${src_col}${PR_L},F{rr}),0)"
widths(bc, {"A": 14, "B": 14})
bc.sheet_properties.tabColor = "A6A6A6"

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != F:
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.active = 0
wb.save(out)
print("saved", out)
